"""Runner: one Paper2Poster attempt per (model, paper, attempt), logged to runs.csv.

Implements component 1 of spec/02_architecture.md. Each attempt is a separate
Paper2Poster process; only `llm_model` changes between runs. With
vision_model: same_as_candidate the candidate fills both the text and the
vision (critic) slot.

Usage (from harness/):
    python src/runner.py --dry-run           # 1 paper x all candidates -> data/dry_run/
    python src/runner.py --estimate          # scale dry-run cost/latency to the full run
    python src/runner.py                     # full run -> data/runs.csv, outputs/<run_id>/
    python src/runner.py --models gpt-5.6-sol --papers P01 P02   # subset

Re-running resumes: (model, paper, attempt) triples already in the runs file are
skipped. Failed attempts are kept as data (FR-7) and are not retried on resume.

workflow_success here covers the automatic checks only (E_API, E_PIPELINE,
E_NO_FILE, E_UNREADABLE). The team fills `visual_check` afterwards (OI-7).
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

HARNESS = Path(__file__).resolve().parent.parent

RUN_COLUMNS = [
    "run_id", "run_order", "batch_id", "llm_model", "input_paper_id", "attempt",
    "prompt_sha256", "p2p_commit", "p2p_patch_sha256", "config_sha256", "vision_model", "temperature", "seed",
    "started_at", "finished_at", "latency",
    "input_tokens", "output_tokens", "cached_tokens",
    "vision_input_tokens", "vision_output_tokens", "token_usage",
    "generation_cost", "vision_cost", "price_table_as_of",
    "retries", "workflow_success", "failure_reason", "visual_check",
    "poster_path", "poster_word_count",
]

# Signs of a gateway/transport problem in Paper2Poster's output (OpenAI SDK
# exception names and HTTP codes). Anything else non-zero is E_PIPELINE.
TRANSPORT_ERRORS = re.compile(
    r"APIConnectionError|APITimeoutError|RateLimitError|InternalServerError"
    r"|Error code: (?:429|5\d\d)|ReadTimeout|ConnectTimeout"
    r"|Connection (?:reset|refused|aborted)"
)
PAPER_ID = re.compile(r"^[A-Za-z0-9_-]+$")
BACKOFF_S = 30
SAME_AS_CANDIDATE = "same_as_candidate"


def now_utc() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def resolve(p: str) -> Path:
    path = Path(p)
    return path if path.is_absolute() else HARNESS / path


def vision_for(model: str, p2c: dict) -> str:
    """Vision (critic) model for a candidate: the candidate itself, or one fixed model (OI-1)."""
    v = p2c["vision_model"]
    return model if v == SAME_AS_CANDIDATE else v


def config_sha256(cfg: dict) -> str:
    """Hash of the `fixed` section only, so adding a candidate keeps it stable."""
    return hashlib.sha256(yaml.safe_dump(cfg["fixed"], sort_keys=True).encode()).hexdigest()


def prompt_sha256(p2p: Path) -> str:
    """Hash of Paper2Poster's agent prompts and style defaults (the "fixed prompt")."""
    files = sorted((p2p / "utils" / "prompt_templates").glob("*.yaml")) + [p2p / "config" / "poster.yaml"]
    h = hashlib.sha256()
    for f in files:
        h.update(str(f.relative_to(p2p)).encode())
        h.update(f.read_bytes())
    return h.hexdigest()


def check_p2p_commit(p2p: Path, expected: str, patch: Path | None) -> tuple[str, str]:
    """Return (HEAD, patch sha256) of the Paper2Poster checkout.

    Exit unless it is the pinned commit with exactly the configured patch applied
    (applied here if the checkout is clean), or no changes at all if there is no patch.
    """
    # Pin every git setting that changes `git diff` output, so the comparison with the
    # patch file doesn't depend on the user's git config.
    git = ["git", "-C", str(p2p), "-c", "color.ui=never", "-c", "diff.noprefix=false",
           "-c", "diff.mnemonicPrefix=false", "-c", "diff.algorithm=myers", "-c", "core.autocrlf=false"]
    head = subprocess.run(git + ["rev-parse", "HEAD"], capture_output=True, text=True, check=True).stdout.strip()
    if head != expected:
        sys.exit(f"Paper2Poster is at {head}, config pins {expected}. Run: git -C {p2p} checkout {expected}")
    diff = lambda: subprocess.run(git + ["diff", "--no-ext-diff", "--unified=3", "--inter-hunk-context=0",
                                         "--indent-heuristic", "--src-prefix=a/", "--dst-prefix=b/"],
                                  capture_output=True, text=True, check=True).stdout
    want = patch.read_text() if patch else ""
    if patch and diff() == "":
        subprocess.run(git + ["apply", str(patch)], check=True)
    if diff() != want:
        sys.exit(f"Paper2Poster checkout has local changes other than {patch or 'none'} "
                 f"(scope rule: no pipeline edits). Inspect with: git -C {p2p} diff")
    return head, hashlib.sha256(want.encode()).hexdigest() if patch else ""


def list_papers(papers_dir: Path) -> list[tuple[str, Path]]:
    papers = []
    for pdf in sorted(papers_dir.glob("*.pdf")):
        if not PAPER_ID.match(pdf.stem):
            sys.exit(f"Bad paper file name {pdf.name}: use <input_paper_id>.pdf with letters, digits, _ or -")
        papers.append((pdf.stem, pdf))
    return papers


def read_runs(path: Path) -> list[dict]:
    if not path.exists():
        return []
    with open(path, newline="") as f:
        return list(csv.DictReader(f))


def append_run(path: Path, row: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    new = not path.exists()
    with open(path, "a", newline="") as f:
        w = csv.DictWriter(f, fieldnames=RUN_COLUMNS)
        if new:
            w.writeheader()
        w.writerow({k: "" if row.get(k) is None else row.get(k) for k in RUN_COLUMNS})


def load_prices(path: Path) -> tuple[dict, str]:
    if not path.exists():
        return {}, ""
    p = yaml.safe_load(open(path)) or {}
    return p.get("models") or {}, str(p.get("as_of") or "")


def token_cost(prices: dict, model: str, tin, tout):
    """USD from per-million-token list prices; None if tokens or prices are missing."""
    pr = prices.get(model) or {}
    pin, pout = pr.get("input_per_mtok"), pr.get("output_per_mtok")
    if tin is None or tout is None or pin is None or pout is None:
        return None
    return round((tin * pin + tout * pout) / 1e6, 6)


def run_cost(row: dict) -> float:
    return sum(float(row[k]) for k in ("generation_cost", "vision_cost") if row.get(k) not in (None, ""))


def poster_text(pptx_path: Path) -> str:
    """All text on the poster slide(s), including grouped shapes and tables."""
    def walk(shapes):
        for s in shapes:
            if s.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from walk(s.shapes)
            elif s.has_text_frame:
                yield s.text_frame.text
            elif getattr(s, "has_table", False):
                for r in s.table.rows:
                    for c in r.cells:
                        yield c.text
    prs = Presentation(str(pptx_path))
    return "\n".join(t for slide in prs.slides for t in walk(slide.shapes) if t.strip())


def read_tokens(out_dir: Path) -> dict:
    """Token totals summed from detail_log.json's per-step `<step>_{in,out}_{t,v}` counts.

    log.json's text totals are wrong at the pinned commit: new_pipeline.py adds the
    outline step's tokens again where it should add the content step's, so the
    content step (most of the text tokens) is missing.
    """
    detail = out_dir / "detail_log.json"
    if not detail.exists():
        return {}
    d = json.loads(detail.read_text())
    def total(suffix):
        return sum(v for k, v in d.items() if k.endswith(suffix) and isinstance(v, (int, float)))
    tin, tout = total("_in_t"), total("_out_t")
    return {"input_tokens": tin, "output_tokens": tout,
            "vision_input_tokens": total("_in_v"), "vision_output_tokens": total("_out_v"),
            "token_usage": tin + tout}


def archive_intermediates(p2p: Path, tag: str, paper_id: str, out_dir: Path) -> None:
    """Copy the parsed content, layout, captions and placed figures into out_dir/intermediate/.

    Paper2Poster keys these by model and paper name, so the next attempt of the same
    pair overwrites them. src/latex_poster.py builds the LaTeX poster from this copy.
    """
    inter = out_dir / "intermediate"
    files = {
        f"contents/{tag}_{paper_id}_raw_content.json": "raw_content.json",
        f"contents/{tag}_{paper_id}_bullet_point_content_0.json": "bullet_point_content.json",
        f"tree_splits/{tag}_{paper_id}_tree_split_0.json": "tree_split.json",
        f"{tag}_images_and_tables/{paper_id}_images.json": "images.json",
        f"{tag}_images_and_tables/{paper_id}_tables.json": "tables.json",
    }
    for src, dst in files.items():
        if (p2p / src).exists():
            inter.mkdir(parents=True, exist_ok=True)
            shutil.copyfile(p2p / src, inter / dst)
    if (inter / "tree_split.json").exists():
        for fig in json.loads((inter / "tree_split.json").read_text()).get("figure_arrangement", []):
            rel = fig.get("figure_path")
            if rel and (p2p / rel).exists():
                (inter / rel).parent.mkdir(parents=True, exist_ok=True)
                shutil.copyfile(p2p / rel, inter / rel)


def run_attempt(*, run_id: int, model: str, paper_id: str, pdf: Path, cfg: dict,
                out_root: Path, env: dict) -> dict:
    """Run one Paper2Poster attempt. Never raises on model/pipeline failure."""
    p2c = cfg["fixed"]["paper2poster"]
    p2p = resolve(p2c["path"])
    vision = vision_for(model, p2c)
    max_retries = cfg["gateway"].get("max_transport_retries", 2)

    staged = p2p / "harness_input" / paper_id
    staged.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(pdf, staged / "paper.pdf")
    poster_rel = f"harness_input/{paper_id}/paper.pdf"
    tmp_rel = f"tmp/run_{run_id}"
    # Same expression Paper2Poster uses for its output folder (new_pipeline.py, step 8).
    gen_dir = p2p / f"<{model}_{vision}>_generated_posters" / poster_rel.replace("paper.pdf", "")
    if gen_dir.exists():
        shutil.rmtree(gen_dir)  # left over from an interrupted attempt

    cmd = [str(resolve(p2c["python"])), "-m", "PosterAgent.new_pipeline",
           f"--poster_path={poster_rel}", f"--model_name_t={model}", f"--model_name_v={vision}",
           f"--tmp_dir={tmp_rel}",
           f"--poster_width_inches={p2c['poster_width_inches']}",
           f"--poster_height_inches={p2c['poster_height_inches']}"]
    if p2c.get("institution_logo_path"):
        cmd.append(f"--institution_logo_path={resolve(p2c['institution_logo_path'])}")
    cmd += [str(a) for a in p2c.get("extra_args") or []]

    out_dir = out_root / str(run_id)
    out_dir.mkdir(parents=True, exist_ok=True)
    log_path = out_dir / "run.log"

    row = {"retries": 0, "workflow_success": 0}
    for attempt_try in range(max_retries + 1):
        row["started_at"] = now_utc()
        t0 = time.monotonic()
        with open(log_path, "a") as log:
            log.write(f"=== try {attempt_try + 1} {row['started_at']}\n$ {' '.join(cmd)}\n")
            log.flush()
            try:
                rc = subprocess.run(cmd, cwd=p2p, env=env, stdout=log, stderr=subprocess.STDOUT,
                                    timeout=p2c.get("timeout_s", 3600)).returncode
            except subprocess.TimeoutExpired:
                rc = None
                log.write("\n=== harness: attempt timed out\n")
        row["latency"] = round(time.monotonic() - t0, 2)
        row["finished_at"] = now_utc()
        if rc == 0:
            break
        this_try = log_path.read_text(errors="replace").rsplit("=== try ", 1)[-1]
        transport = rc is not None and bool(TRANSPORT_ERRORS.search(this_try))
        if transport and attempt_try < max_retries:
            row["retries"] += 1
            time.sleep(BACKOFF_S * 2 ** attempt_try)
            continue
        row["failure_reason"] = "E_API" if transport else "E_PIPELINE"
        break

    shutil.rmtree(p2p / tmp_rel, ignore_errors=True)
    if gen_dir.exists():
        for item in gen_dir.iterdir():
            shutil.move(str(item), str(out_dir / item.name))
        shutil.rmtree(gen_dir, ignore_errors=True)
    archive_intermediates(p2p, f"<{model}_{vision}>", paper_id, out_dir)

    row.update(read_tokens(out_dir))

    if row.get("failure_reason"):
        return row

    pptx = out_dir / f"{paper_id}.pptx"
    if not pptx.exists():
        row["failure_reason"] = "E_NO_FILE"
        return row
    try:
        text = poster_text(pptx)
    except Exception as e:  # corrupt or unreadable PPTX
        with open(log_path, "a") as log:
            log.write(f"\n=== harness: text extraction failed: {e!r}\n")
        text = ""
    (out_dir / "poster_text.txt").write_text(text)
    words = len(text.split())
    if words == 0:
        row["failure_reason"] = "E_UNREADABLE"
        return row

    row.update(workflow_success=1, poster_path=str(pptx.relative_to(HARNESS)), poster_word_count=words)
    return row


def estimate(cfg: dict) -> None:
    """Scale mean dry-run cost and latency per model to the full run (NFR-5)."""
    rows = [r for r in read_runs(HARNESS / "data" / "dry_run" / "runs.csv") if r["workflow_success"] == "1"]
    if not rows:
        sys.exit("No successful dry-run attempts in data/dry_run/runs.csv. Run --dry-run first.")
    n = cfg["fixed"]["n_papers"] * cfg["fixed"].get("attempts_per_pair", 1)
    total, missing = 0.0, []
    print(f"{'model':40} {'cost/poster':>12} {'latency/poster':>15} {'full run cost':>14}")
    for c in cfg["candidates"]:
        mr = [r for r in rows if r["llm_model"] == c["id"]]
        if not mr:
            print(f"{c['id']:40} {'no successful dry run':>43}")
            missing.append(c["id"])
            continue
        if any(r["generation_cost"] == "" for r in mr):
            missing.append(c["id"])
        cost = sum(run_cost(r) for r in mr) / len(mr)
        lat = sum(float(r["latency"]) for r in mr) / len(mr)
        total += cost * n
        print(f"{c['id']:40} {cost:12.4f} {lat / 60:12.1f} min {cost * n:14.2f}")
    print(f"\nEstimated full-run cost ({n} attempts per model): ${total:.2f}")
    if missing:
        print(f"Incomplete (no dry run or no prices in prices.yaml): {', '.join(missing)}")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--config", default="config/experiment.yaml")
    ap.add_argument("--dry-run", action="store_true", help="first paper only; results go to data/dry_run/")
    ap.add_argument("--estimate", action="store_true", help="estimate full-run cost from the dry run, then exit")
    ap.add_argument("--models", nargs="+", help="subset of candidate IDs")
    ap.add_argument("--papers", nargs="+", help="subset of input_paper_ids")
    args = ap.parse_args()

    cfg = yaml.safe_load(open(resolve(args.config)))
    if args.estimate:
        return estimate(cfg)

    fixed, p2c, g = cfg["fixed"], cfg["fixed"]["paper2poster"], cfg["gateway"]
    if p2c.get("vision_model") in (None, "", "TBD"):
        sys.exit(f"Set fixed.paper2poster.vision_model in the config: a gateway ID or {SAME_AS_CANDIDATE} (OI-1).")
    logo = p2c.get("institution_logo_path")
    if logo == "TBD":
        sys.exit("Set fixed.paper2poster.institution_logo_path to a logo file, or null to let Paper2Poster find one (OI-18).")
    if logo and not resolve(logo).is_file():
        sys.exit(f"Logo file not found: {resolve(logo)}")
    base_url, api_key = os.environ.get(g["base_url_env"]), os.environ.get(g["api_key_env"])
    if not base_url or not api_key:
        sys.exit(f"Set {g['base_url_env']} and {g['api_key_env']} first.")
    # camel's OpenAICompatibleModel reads the COMPATIBILIY variables (its spelling) but then
    # builds the client from the unset constructor args, so the OpenAI SDK falls back to
    # OPENAI_API_KEY / OPENAI_BASE_URL. Set both pairs.
    env = dict(os.environ,
               OPENAI_COMPATIBILIY_API_KEY=api_key,
               OPENAI_COMPATIBILIY_API_BASE_URL=base_url,
               OPENAI_API_KEY=api_key,
               OPENAI_BASE_URL=base_url)
    # First-run model downloads (docling, EasyOCR): Hugging Face's Xet transfer timed out
    # on slow networks; plain HTTP with a longer timeout didn't. Users can still override.
    env.setdefault("HF_HUB_DISABLE_XET", "1")
    env.setdefault("HF_HUB_DOWNLOAD_TIMEOUT", "120")
    if p2c.get("max_output_tokens"):
        if not p2c.get("patch"):
            sys.exit("fixed.paper2poster.max_output_tokens needs the max-output-tokens patch (OI-17).")
        env["P2P_MAX_OUTPUT_TOKENS"] = str(p2c["max_output_tokens"])  # read by the patch

    p2p = resolve(p2c["path"])
    commit, patch_sha = check_p2p_commit(p2p, p2c["commit"], resolve(p2c["patch"]) if p2c.get("patch") else None)
    hashes = {"prompt_sha256": prompt_sha256(p2p), "p2p_commit": commit, "p2p_patch_sha256": patch_sha,
              "config_sha256": config_sha256(cfg)}

    papers = list_papers(resolve(fixed["papers_dir"]))
    if not papers:
        sys.exit(f"No PDFs in {fixed['papers_dir']}")
    if not args.dry_run and not args.papers and len(papers) != fixed["n_papers"]:
        sys.exit(f"Found {len(papers)} papers, config expects {fixed['n_papers']} (DR-7).")
    order = {pid: i + 1 for i, (pid, _) in enumerate(papers)}  # run_order, same for every model
    if args.papers:
        unknown = set(args.papers) - set(order)
        if unknown:
            sys.exit(f"Unknown paper IDs: {sorted(unknown)}")
        papers = [p for p in papers if p[0] in args.papers]
    if args.dry_run:
        papers = papers[:1]

    models = [c["id"] for c in cfg["candidates"]]
    if args.models:
        unknown = set(args.models) - set(models)
        if unknown:
            sys.exit(f"Not in config candidates: {sorted(unknown)}")
        models = [m for m in models if m in args.models]

    data_dir = HARNESS / "data" / ("dry_run" if args.dry_run else "")
    runs_path = data_dir / "runs.csv"
    out_root = HARNESS / "outputs" / ("dry_run" if args.dry_run else "")
    prices, as_of = load_prices(resolve(fixed.get("prices_file", "config/prices.yaml")))
    cap = cfg["experiment"].get("spend_cap_usd")
    batch_size = cfg.get("spc", {}).get("p_chart_batch_size", 5)

    runs = read_runs(runs_path)
    for r in runs:
        if (r["p2p_commit"] != commit or r["prompt_sha256"] != hashes["prompt_sha256"]
                or r.get("p2p_patch_sha256", "") != patch_sha):
            sys.exit(f"{runs_path} has run {r['run_id']} from a different Paper2Poster commit, patch or prompts (DR-6).")
    done = {(r["llm_model"], r["input_paper_id"], int(r["attempt"])) for r in runs}
    next_id = max((int(r["run_id"]) for r in runs), default=0) + 1
    spent = sum(run_cost(r) for r in runs)
    if cap is not None and any(not (prices.get(m) or {}).get("input_per_mtok") for m in models):
        print("WARNING: spend_cap_usd is set but prices.yaml is incomplete; the cap can't count those models.")

    for model in models:
        for paper_id, pdf in papers:
            for attempt in range(1, fixed.get("attempts_per_pair", 1) + 1):
                if (model, paper_id, attempt) in done:
                    continue
                if cap is not None and spent >= cap:
                    sys.exit(f"Spend cap reached: ${spent:.2f} >= ${cap} (NFR-5). Stopping.")
                vision = vision_for(model, p2c)
                print(f"[run {next_id}] {model} × {paper_id} (attempt {attempt}) ...", flush=True)
                row = run_attempt(run_id=next_id, model=model, paper_id=paper_id, pdf=pdf,
                                  cfg=cfg, out_root=out_root, env=env)
                row.update(hashes, run_id=next_id, run_order=order[paper_id],
                           batch_id=math.ceil(order[paper_id] / batch_size),
                           llm_model=model, input_paper_id=paper_id, attempt=attempt,
                           vision_model=vision,
                           temperature=fixed.get("generation", {}).get("temperature"),
                           generation_cost=token_cost(prices, model, row.get("input_tokens"), row.get("output_tokens")),
                           vision_cost=token_cost(prices, vision,
                                                  row.get("vision_input_tokens"), row.get("vision_output_tokens")),
                           price_table_as_of=as_of)
                append_run(runs_path, row)
                spent += run_cost(row)
                status = "ok" if row["workflow_success"] else row["failure_reason"]
                print(f"[run {next_id}] {status}  {row['latency']}s  tokens={row.get('token_usage')}", flush=True)
                next_id += 1


if __name__ == "__main__":
    main()
